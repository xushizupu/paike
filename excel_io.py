"""把排课结果导出为三个工作簿：班级课表、教师课表、校验统计。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

import scheduler


THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_HEADER = Font(name="微软雅黑", size=11, bold=True)
FONT_TITLE = Font(name="微软雅黑", size=14, bold=True)
FONT_BODY = Font(name="微软雅黑", size=10)
ROW_HEIGHT = 39


def _set_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_header_row(ws, row, values):
    ws.row_dimensions[row].height = 24
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = FONT_HEADER
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_simple_cell(ws, row, col, value, font=FONT_BODY):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def _write_rich_cell(ws, row, col, primary, secondary):
    if not primary:
        return _write_simple_cell(ws, row, col, "")
    rich = CellRichText(
        TextBlock(InlineFont(rFont="微软雅黑", sz=12, b=True), primary),
        TextBlock(InlineFont(rFont="微软雅黑", sz=9, b=False), "\n" + secondary),
    )
    cell = ws.cell(row=row, column=col, value=rich)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def _write_title(ws, row, title, total_columns):
    ws.row_dimensions[row].height = 30
    ws.cell(row=row, column=1, value=title).font = FONT_TITLE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="center")


def _build_class_sheet(ws, cls, data, settings, result):
    total_columns = 3 + len(data.days)
    _set_widths(ws, [12, 8, 15] + [17] * len(data.days))
    title_suffix = settings["titles"].get("classSuffix") or "课程表"
    _write_title(ws, 1, f"{cls}{title_suffix}", total_columns)
    _write_header_row(ws, 2, ["时段", "节次", "时间"] + data.days)
    row = 3
    section_merges = []
    for period_index, period in enumerate(data.periods):
        ws.row_dimensions[row].height = ROW_HEIGHT
        span = data.section_spans[period_index] if period_index < len(data.section_spans) else 0
        if span > 0:
            _write_simple_cell(ws, row, 1, period["section"] or "", FONT_HEADER)
            section_merges.append((row, span))
        _write_simple_cell(ws, row, 2, period["period"], FONT_HEADER)
        _write_simple_cell(ws, row, 3, period["time"], FONT_HEADER)
        for day_idx in range(len(data.days)):
            key = scheduler.slot_key(day_idx, period_index)
            if key in settings["classBlocked"].get(cls, set()):
                _write_simple_cell(ws, row, 4 + day_idx, "")
            else:
                cell = result.get(cls, {}).get(key)
                if cell:
                    _write_rich_cell(ws, row, 4 + day_idx, cell["subject"], cell["teacher"])
                else:
                    _write_simple_cell(ws, row, 4 + day_idx, "")
        row += 1
    for start_row, span in section_merges:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + span - 1, end_column=1)


def _build_teacher_sheet(ws, teacher, data, settings, result):
    total_columns = 3 + len(data.days)
    _set_widths(ws, [12, 8, 15] + [18] * len(data.days))
    title_suffix = settings["titles"].get("teacherSuffix") or "课程表"
    _write_title(ws, 1, f"{teacher}{title_suffix}", total_columns)
    _write_header_row(ws, 2, ["时段", "节次", "时间"] + data.days)
    row = 3
    section_merges = []
    for period_index, period in enumerate(data.periods):
        ws.row_dimensions[row].height = ROW_HEIGHT
        span = data.section_spans[period_index] if period_index < len(data.section_spans) else 0
        if span > 0:
            _write_simple_cell(ws, row, 1, period["section"] or "", FONT_HEADER)
            section_merges.append((row, span))
        _write_simple_cell(ws, row, 2, period["period"], FONT_HEADER)
        _write_simple_cell(ws, row, 3, period["time"], FONT_HEADER)
        for day_idx in range(len(data.days)):
            key = scheduler.slot_key(day_idx, period_index)
            entries = [
                (cell["subject"], f"{cls}")
                for cls in data.classes
                if (cell := result.get(cls, {}).get(key)) and cell.get("teacher") == teacher
            ]
            if key in settings["teacherUnavailable"].get(teacher, set()) or not entries:
                _write_simple_cell(ws, row, 4 + day_idx, "")
            elif len(entries) == 1:
                subject, cls_name = entries[0]
                _write_rich_cell(ws, row, 4 + day_idx, subject, cls_name)
            else:
                _write_rich_cell(ws, row, 4 + day_idx, entries[0][0], "\n".join(item[1] for item in entries))
        row += 1
    for start_row, span in section_merges:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + span - 1, end_column=1)


def _build_report_sheet(ws, data, settings, result, issues):
    _set_widths(ws, [18, 100])
    ws.cell(row=1, column=1, value="校验报告").font = FONT_TITLE
    row = 3
    summary = [
        f"班级数量：{len(data.classes)}",
        f"科目数量：{len(data.subjects)}",
        f"每天节次：{data.n_slots}",
        f"每周可用格：{len(data.days) * data.n_slots}",
        f"校验问题数：{len(issues)}",
    ]
    for line in summary:
        _write_simple_cell(ws, row, 1, line)
        row += 1
    row += 1
    if issues:
        _write_simple_cell(ws, row, 1, "问题明细", FONT_HEADER)
        row += 1
        for issue in issues:
            _write_simple_cell(ws, row, 1, "×")
            _write_simple_cell(ws, row, 2, issue)
            row += 1
    else:
        _write_simple_cell(ws, row, 1, "未发现问题")


def _build_stats_sheet(ws, data, result):
    _set_widths(ws, [12, 16, 12, 10, 10, 12])
    _write_header_row(ws, 1, ["班级", "科目", "教师", "应排", "实排", "状态"])
    row = 2
    for cls in data.classes:
        for subject in data.subjects:
            required = data.hours[(cls, subject)]
            actual = sum(1 for cell in result.get(cls, {}).values() if cell.get("subject") == subject)
            teacher = data.teachers.get((cls, subject), "")
            status = "正确" if actual == required else "不一致"
            for col, value in enumerate([cls, subject, teacher, required, actual, status], start=1):
                _write_simple_cell(ws, row, col, value)
            row += 1


def _build_class_workbook(data, settings, result):
    wb = Workbook()
    wb.remove(wb.active)
    for cls in data.classes:
        ws = wb.create_sheet(title=cls[:31])
        _build_class_sheet(ws, cls, data, settings, result)
    return wb


def _build_teacher_workbook(data, settings, result):
    wb = Workbook()
    wb.remove(wb.active)
    for teacher in data.teacher_list:
        ws = wb.create_sheet(title=teacher[:31])
        _build_teacher_sheet(ws, teacher, data, settings, result)
    return wb


def _build_check_workbook(data, settings, result, issues):
    wb = Workbook()
    ws = wb.active
    ws.title = "校验报告"
    _build_report_sheet(ws, data, settings, result, issues)
    ws2 = wb.create_sheet("课时统计")
    _build_stats_sheet(ws2, data, result)
    return wb


def write_result_files(data, settings, result, output_dir, timestamp, suffix):
    """导出三个工作簿并返回文件信息列表。"""
    settings = scheduler.normalize_settings(settings, data)
    issues = scheduler.validate_result(data, settings, result)

    class_wb = _build_class_workbook(data, settings, result)
    teacher_wb = _build_teacher_workbook(data, settings, result)
    check_wb = _build_check_workbook(data, settings, result, issues)

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    files = []
    for wb, name in [
        (class_wb, f"班级课程表_{timestamp}_{suffix}.xlsx"),
        (teacher_wb, f"教师课程表_{timestamp}_{suffix}.xlsx"),
        (check_wb, f"排课校验_{timestamp}_{suffix}.xlsx"),
    ]:
        path = output_dir / name
        wb.save(path)
        files.append({"name": name, "path": str(path)})
    return issues, files
